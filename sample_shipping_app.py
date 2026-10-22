from __future__ import annotations

import os
from io import BytesIO
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from address_utils import normalize_text, split_address, validate_address_groups, validate_batch_rows
from excel_utils import (
    build_batch_template_file,
    build_output_file,
    load_product_options,
    parse_uploaded_address_file,
    save_output_file,
)
from github_sync import replace_template
from time_utils import now_shanghai
from ui_components import render_address_preview, render_product_table


BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = BASE_DIR / "样品寄送模版.xlsx"
# 判断是否跑在 Streamlit Community Cloud 上：云端目录在 /mount/src，且 HOME 为 /home/appuser
IS_CLOUD = str(BASE_DIR).startswith("/mount/src") or os.environ.get("HOME", "").rstrip("/").endswith("appuser")
GITHUB_REPO = "xiaofengxu-art/sample-shipping-app"
TEMPLATE_REPO_PATH = "样品寄送模版.xlsx"
def get_secret(name: str) -> str | None:
    """读取 Streamlit secrets，未配置时返回 None 而不是抛异常。"""
    try:
        value = st.secrets[name]
    except Exception:
        return None
    return str(value) if value else None


def validate_template_bytes(data: bytes) -> tuple[bool, str]:
    """检查上传的模板是否可用，返回 (是否通过, 说明)。"""
    try:
        workbook = load_workbook(BytesIO(data), data_only=False)
    except Exception:
        return False, "无法解析该文件，请确认是 .xlsx 格式。"
    missing = [name for name in ("Sheet0", "Sheet1") if name not in workbook.sheetnames]
    if missing:
        return False, "模板缺少工作表：" + "、".join(missing) + "。"
    rows = 0
    for row in workbook["Sheet1"].iter_rows(min_row=2, values_only=True):
        if row and row[0] not in (None, ""):
            rows += 1
    if rows == 0:
        return False, "Sheet1 里没有读到商品数据。"
    return True, f"Sheet1 共 {rows} 行商品"


CUSTOMER_OPTIONS = [
    {"code": "KH00832", "name": "第三方样品"},
    {"code": "KH00007", "name": "小红书"},
]


def address_order_group_count_key(address_index: int) -> str:
    return f"address_order_group_count_{address_index}"


def group_item_count_key(address_index: int, group_index: int) -> str:
    return f"group_item_count_{address_index}_{group_index}"


def group_item_count_input_key(address_index: int, group_index: int) -> str:
    return f"group_item_count_input_{address_index}_{group_index}"


def increment_group_item_count(address_index: int, group_index: int) -> None:
    item_key = group_item_count_key(address_index, group_index)
    item_input_key = group_item_count_input_key(address_index, group_index)
    new_value = max(1, int(st.session_state.get(item_key, 1))) + 1
    st.session_state[item_key] = new_value
    st.session_state[item_input_key] = new_value


def decrement_group_item_count(address_index: int, group_index: int) -> None:
    item_key = group_item_count_key(address_index, group_index)
    item_input_key = group_item_count_input_key(address_index, group_index)
    new_value = max(1, int(st.session_state.get(item_key, 1)) - 1)
    st.session_state[item_key] = new_value
    st.session_state[item_input_key] = new_value


def ensure_address_defaults(address_index: int) -> None:
    order_count_key = address_order_group_count_key(address_index)
    if order_count_key not in st.session_state:
        st.session_state[order_count_key] = 1
    st.session_state[order_count_key] = max(1, int(st.session_state[order_count_key]))

    for group_index in range(int(st.session_state[order_count_key])):
        item_key = group_item_count_key(address_index, group_index)
        item_input_key = group_item_count_input_key(address_index, group_index)
        if item_key not in st.session_state:
            st.session_state[item_key] = 1
        st.session_state[item_key] = max(1, int(st.session_state[item_key]))
        if item_input_key not in st.session_state:
            st.session_state[item_input_key] = st.session_state[item_key]
        st.session_state[item_input_key] = max(1, int(st.session_state[item_input_key]))


def copy_previous_address_items(address_index: int) -> None:
    if address_index <= 0:
        return

    previous_order_count_key = address_order_group_count_key(address_index - 1)
    current_order_count_key = address_order_group_count_key(address_index)
    previous_order_count = int(st.session_state.get(previous_order_count_key, 1))
    st.session_state[current_order_count_key] = previous_order_count

    for group_index in range(previous_order_count):
        prev_item_count_key = group_item_count_key(address_index - 1, group_index)
        curr_item_count_key = group_item_count_key(address_index, group_index)
        previous_item_count = int(st.session_state.get(prev_item_count_key, 1))
        st.session_state[curr_item_count_key] = previous_item_count
        st.session_state[group_item_count_input_key(address_index, group_index)] = previous_item_count
        st.session_state[f"external_order_no_{address_index}_{group_index}"] = ""

        for item_index in range(previous_item_count):
            for prefix in ("product", "quantity"):
                prev_key = f"{prefix}_{address_index - 1}_{group_index}_{item_index}"
                curr_key = f"{prefix}_{address_index}_{group_index}_{item_index}"
                if prev_key in st.session_state:
                    st.session_state[curr_key] = st.session_state[prev_key]


def build_batch_preview_rows(batch_rows: list[dict[str, str]]) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    preview_rows: list[dict[str, str]] = []
    failed_rows: list[dict[str, str]] = []
    for row in batch_rows:
        try:
            parsed = split_address(row["address_text"])
            preview_rows.append(
                {
                    "外部平台单号": row["external_order_no"],
                    "收货人": parsed["receiver"],
                    "手机": parsed["phone"],
                    "省": parsed["province"],
                    "市": parsed["city"],
                    "区": parsed["district"],
                    "详细地址": parsed["detail_address"],
                    "状态": "解析成功",
                }
            )
        except ValueError as exc:
            preview_rows.append(
                {
                    "外部平台单号": row["external_order_no"],
                    "收货人": "",
                    "手机": "",
                    "省": "",
                    "市": "",
                    "区": "",
                    "详细地址": f"解析失败：{exc}",
                    "状态": "解析失败",
                }
            )
            failed_rows.append(
                {
                    "外部平台单号": row["external_order_no"],
                    "地址": row["address_text"],
                    "失败原因": str(exc),
                }
            )
    return preview_rows, failed_rows


def collect_cross_source_duplicate_errors(
    active_manual_groups: list[dict[str, object]],
    batch_rows: list[dict[str, str]],
) -> list[str]:
    combined_seen_orders: dict[str, str] = {}
    duplicate_errors: list[str] = []

    for address_index, address_group in enumerate(active_manual_groups, start=1):
        for order_index, order_group in enumerate(address_group["order_groups"], start=1):
            order_no = normalize_text(order_group["external_order_no"])
            if not order_no:
                continue
            if order_no in combined_seen_orders:
                duplicate_errors.append(
                    f"外部平台单号 {order_no} 重复出现：{combined_seen_orders[order_no]} 与 地址 {address_index} 的外部单号 {order_index}。"
                )
            else:
                combined_seen_orders[order_no] = f"地址 {address_index} 的外部单号 {order_index}"

    for row_index, row in enumerate(batch_rows, start=1):
        order_no = normalize_text(row["external_order_no"])
        if not order_no:
            continue
        if order_no in combined_seen_orders:
            duplicate_errors.append(
                f"外部平台单号 {order_no} 重复出现：{combined_seen_orders[order_no]} 与 批量地址第 {row_index} 行。"
            )
        else:
            combined_seen_orders[order_no] = f"批量地址第 {row_index} 行"

    return duplicate_errors


def sync_customer_from_code() -> None:
    selected_code = st.session_state.customer_code_select
    matched = next(option for option in CUSTOMER_OPTIONS if option["code"] == selected_code)
    st.session_state.customer_name_select = matched["name"]


def sync_customer_from_name() -> None:
    selected_name = st.session_state.customer_name_select
    matched = next(option for option in CUSTOMER_OPTIONS if option["name"] == selected_name)
    st.session_state.customer_code_select = matched["code"]


st.set_page_config(page_title="样品寄送单生成器", page_icon="📦", layout="centered")
st.title("样品寄送单生成器")
st.caption("读取模板中的商品资料，自动拆分地址并生成可下载的样品寄送 Excel。")

if "template_flash" in st.session_state:
    level, text = st.session_state.pop("template_flash")
    getattr(st, level)(text)

if not TEMPLATE_PATH.exists():
    st.error(f"未找到模板文件：{TEMPLATE_PATH}")
    st.stop()

if "product_refresh_time" not in st.session_state:
    st.session_state.product_refresh_time = now_shanghai()
if "address_count" not in st.session_state:
    st.session_state.address_count = 1
if "batch_item_count" not in st.session_state:
    st.session_state.batch_item_count = 1
if "batch_item_count_input" not in st.session_state:
    st.session_state.batch_item_count_input = 1
if "customer_code_select" not in st.session_state:
    st.session_state.customer_code_select = CUSTOMER_OPTIONS[0]["code"]
if "customer_name_select" not in st.session_state:
    st.session_state.customer_name_select = CUSTOMER_OPTIONS[0]["name"]
ensure_address_defaults(0)

st.markdown("### 填写寄送信息")
customer_cols = st.columns(2)
with customer_cols[0]:
    st.selectbox(
        "客户code",
        [option["code"] for option in CUSTOMER_OPTIONS],
        key="customer_code_select",
        on_change=sync_customer_from_code,
    )
with customer_cols[1]:
    st.selectbox(
        "客户名称",
        [option["name"] for option in CUSTOMER_OPTIONS],
        key="customer_name_select",
        on_change=sync_customer_from_name,
    )
selected_customer = next(
    option for option in CUSTOMER_OPTIONS if option["code"] == st.session_state.customer_code_select
)
st.info(
    f"固定写入内容：客户code = {selected_customer['code']}，客户名称 = {selected_customer['name']}，子订单类型 = 样品。支付日期会写入生成当天的静态值。"
)
st.caption("地址输入格式：收货人名称 手机号 详细地址（包含：省市区信息）")

toolbar_col1, toolbar_col2 = st.columns([1, 3])
with toolbar_col1:
    if st.button("刷新商品数据", use_container_width=True):
        load_product_options.clear()
        st.session_state.product_refresh_time = now_shanghai()
        st.success("已重新读取 Sheet1 商品数据。")
with toolbar_col2:
    refresh_time_text = st.session_state.product_refresh_time.strftime("%Y-%m-%d %H:%M:%S")
    st.caption(f"如果你刚更新了模板里的 Sheet1，请先点一次“刷新商品数据”。最近刷新时间：{refresh_time_text}（北京时间）")

products = load_product_options(str(TEMPLATE_PATH))
if not products:
    st.error("未从 Sheet1 读取到商品数据。")
    st.stop()

admin_mode = (not IS_CLOUD) or st.query_params.get("admin") == "1"
if admin_mode:
    with st.expander("更新商品模板（管理员）", expanded=IS_CLOUD):
        st.caption(
            "上传新的 样品寄送模版.xlsx 直接替换当前模板。"
            "本地运行会替换项目里的模板文件；线上会把文件提交回 GitHub 仓库，"
            "Streamlit Cloud 收到提交后自动重新部署，约 1 分钟生效。"
        )
        admin_password = get_secret("ADMIN_PASSWORD")
        upload_token = get_secret("GITHUB_TOKEN")

        allowed = True
        if IS_CLOUD and not upload_token:
            st.warning(
                "线上还没有配置 GITHUB_TOKEN，上传入口不可用。"
                "请在 Streamlit Cloud 的 Settings → Secrets 里补上 GITHUB_TOKEN 和 ADMIN_PASSWORD。"
            )
            allowed = False
        elif IS_CLOUD and not admin_password:
            st.warning("线上模式必须先在 Streamlit 的 secrets 里配置 ADMIN_PASSWORD，上传入口才会开放。")
            allowed = False
        elif IS_CLOUD and not st.session_state.get("admin_ok"):
            st.caption("先输入管理密码，验证通过后才能上传。")
            password_input = st.text_input("管理密码", type="password", key="template_admin_pwd")
            if st.button("确认密码", key="admin_login_button"):
                if password_input == admin_password:
                    st.session_state["admin_ok"] = True
                    st.rerun()
                st.error("管理密码不正确。")
            allowed = False

        if allowed:
            uploaded_template = st.file_uploader(
                "选择新的模板文件", type=["xlsx"], key="template_upload"
            )
            if st.button("替换模板", key="replace_template_button", disabled=uploaded_template is None):
                data = uploaded_template.getvalue() if uploaded_template is not None else b""
                ok, message = validate_template_bytes(data)
                if not ok:
                    st.session_state["template_flash"] = ("error", message)
                    st.rerun()

                if upload_token:
                    try:
                        commit_sha = replace_template(GITHUB_REPO, TEMPLATE_REPO_PATH, data, upload_token)
                    except RuntimeError as exc:
                        st.session_state["template_flash"] = ("error", f"提交到 GitHub 失败：{exc}")
                    else:
                        st.session_state["template_flash"] = (
                            "success",
                            f"{message}，已提交到仓库（{commit_sha[:7]}）。"
                            "Streamlit Cloud 会自动重新部署，约 1 分钟后模板生效。",
                        )
                    st.rerun()
                else:
                    TEMPLATE_PATH.write_bytes(data)
                    load_product_options.clear()
                    st.session_state.product_refresh_time = now_shanghai()
                    st.session_state["template_flash"] = ("success", f"{message}，已替换本地模板文件。")
                    st.rerun()

st.markdown("### 地址与商品明细")
address_toolbar_col1, address_toolbar_col2 = st.columns([1, 1])
with address_toolbar_col1:
    if st.button("新增一个地址", use_container_width=True):
        new_address_index = st.session_state.address_count
        st.session_state.address_count += 1
        ensure_address_defaults(new_address_index)
with address_toolbar_col2:
    if st.button("删除最后一个地址", use_container_width=True, disabled=st.session_state.address_count <= 1):
        st.session_state.address_count -= 1

address_groups: list[dict[str, object]] = []
for address_index in range(st.session_state.address_count):
    ensure_address_defaults(address_index)
    order_count_key = address_order_group_count_key(address_index)

    st.markdown(f"## 地址 {address_index + 1}")
    address_action_col1, address_action_col2, address_action_col3 = st.columns([1, 1, 2])
    with address_action_col1:
        if st.button("新增外部单号", key=f"add_order_group_{address_index}", use_container_width=True):
            new_group_index = int(st.session_state[order_count_key])
            st.session_state[order_count_key] += 1
            st.session_state[group_item_count_key(address_index, new_group_index)] = 1
    with address_action_col2:
        if st.button(
            "删除外部单号",
            key=f"remove_order_group_{address_index}",
            use_container_width=True,
            disabled=st.session_state[order_count_key] <= 1,
        ):
            st.session_state[order_count_key] -= 1
    with address_action_col3:
        if st.button(
            "复制上一地址商品明细",
            key=f"copy_prev_address_{address_index}",
            use_container_width=True,
            disabled=address_index == 0,
        ):
            copy_previous_address_items(address_index)
            st.success(f"地址 {address_index + 1} 已复制上一地址的商品明细。")

    address_text = st.text_area(
        f"一整段地址 {address_index + 1}",
        height=100,
        key=f"address_text_{address_index}",
        placeholder="例如：收货人名称 手机号 详细地址（包含：省市区信息）",
        help="格式建议：收货人名称 + 手机号 + 详细地址（包含：省市区信息）。系统会自动拆分，支持直辖市。",
    )
    render_address_preview(address_text, address_index)

    order_groups: list[dict[str, object]] = []
    for group_index in range(int(st.session_state[order_count_key])):
        item_count_key = group_item_count_key(address_index, group_index)
        item_count_input_key = group_item_count_input_key(address_index, group_index)
        st.session_state[item_count_key] = max(1, int(st.session_state.get(item_count_key, 1)))
        st.session_state[item_count_input_key] = max(1, int(st.session_state.get(item_count_input_key, st.session_state[item_count_key])))

        with st.container(border=True):
            st.markdown(f"### 外部单号 {group_index + 1}")
            order_header_col1, order_header_col2, order_header_col3 = st.columns([3, 1, 1])
            with order_header_col1:
                external_order_no = st.text_input(
                    f"外部平台单号 {address_index + 1}-{group_index + 1}",
                    key=f"external_order_no_{address_index}_{group_index}",
                    placeholder="请输入平台订单号",
                )
            with order_header_col2:
                item_count = st.number_input(
                    f"商品行数 {address_index + 1}-{group_index + 1}",
                    min_value=1,
                    step=1,
                    key=item_count_input_key,
                )
                st.session_state[item_count_key] = int(item_count)
            with order_header_col3:
                button_col1, button_col2 = st.columns(2)
                with button_col1:
                    st.button(
                        "新增商品",
                        key=f"add_item_{address_index}_{group_index}",
                        use_container_width=True,
                        on_click=increment_group_item_count,
                        args=(address_index, group_index),
                    )
                with button_col2:
                    st.button(
                        "删除商品",
                        key=f"remove_item_{address_index}_{group_index}",
                        use_container_width=True,
                        disabled=st.session_state[item_count_key] <= 1,
                        on_click=decrement_group_item_count,
                        args=(address_index, group_index),
                    )

            group_line_items = render_product_table(
                products=products,
                key_prefix=f"{address_index}_{group_index}",
                row_count=int(st.session_state[item_count_key]),
                title="商品明细表",
            )

        order_groups.append({"external_order_no": external_order_no, "line_items": group_line_items})

    address_groups.append({"address_text": address_text, "order_groups": order_groups})

st.markdown("### 批量上传多地址（同商品）")
st.caption("上传表头为“外部平台单号”“地址”的 .xlsx，下面维护一套公共商品，生成时会应用到上传的每一条地址。")
st.download_button(
    "下载批量上传示例 XLSX",
    data=build_batch_template_file(),
    file_name="批量地址上传示例.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)
uploaded_batch_file = st.file_uploader("上传批量地址文件", type=["xlsx"], accept_multiple_files=False)

batch_rows: list[dict[str, str]] = []
batch_file_error = ""
batch_preview_rows: list[dict[str, str]] = []
batch_preview_failed_rows: list[dict[str, str]] = []
if uploaded_batch_file is not None:
    try:
        batch_rows = parse_uploaded_address_file(uploaded_batch_file)
        st.success(f"已读取批量地址 {len(batch_rows)} 条。")
        batch_preview_rows, batch_preview_failed_rows = build_batch_preview_rows(batch_rows)
    except ValueError as exc:
        batch_file_error = str(exc)
        st.error(batch_file_error)

if batch_preview_rows:
    st.markdown("#### 批量地址解析预览")
    preview_success_count = len(batch_preview_rows) - len(batch_preview_failed_rows)
    preview_error_count = len(batch_preview_failed_rows)
    preview_metric_cols = st.columns(3)
    preview_metric_cols[0].metric("总条数", len(batch_preview_rows))
    preview_metric_cols[1].metric("解析成功", preview_success_count)
    preview_metric_cols[2].metric("解析失败", preview_error_count)
    st.dataframe(pd.DataFrame(batch_preview_rows), use_container_width=True, hide_index=True)
    if batch_preview_failed_rows:
        st.warning("以下批量地址解析失败，生成时会被拦截，请先修正原表后重新上传。")
        st.dataframe(pd.DataFrame(batch_preview_failed_rows), use_container_width=True, hide_index=True)

batch_toolbar_col1, batch_toolbar_col2, batch_toolbar_col3 = st.columns([1, 1, 1.2])
with batch_toolbar_col1:
    if st.button("新增公共商品", use_container_width=True):
        st.session_state.batch_item_count += 1
        st.session_state.batch_item_count_input = st.session_state.batch_item_count
with batch_toolbar_col2:
    if st.button("删除公共商品", use_container_width=True, disabled=st.session_state.batch_item_count <= 1):
        st.session_state.batch_item_count -= 1
        st.session_state.batch_item_count_input = st.session_state.batch_item_count
with batch_toolbar_col3:
    st.session_state.batch_item_count = max(1, int(st.session_state.batch_item_count))
    st.session_state.batch_item_count_input = max(1, int(st.session_state.batch_item_count_input))
    batch_count = st.number_input("公共商品行数", min_value=1, step=1, key="batch_item_count_input")
    st.session_state.batch_item_count = int(batch_count)

batch_shared_line_items = render_product_table(
    products=products,
    key_prefix="batch",
    row_count=st.session_state.batch_item_count,
    title="公共商品明细表",
)

submitted = st.button("生成 Excel", type="primary")

if submitted:
    active_manual_groups = [
        group
        for group in address_groups
        if normalize_text(group["address_text"])
        or any(normalize_text(order_group["external_order_no"]) for order_group in group["order_groups"])
    ]
    manual_validation_errors = validate_address_groups(active_manual_groups)
    batch_validation_errors = validate_batch_rows(batch_rows, batch_shared_line_items) if batch_rows else []
    cross_source_duplicate_errors = collect_cross_source_duplicate_errors(active_manual_groups, batch_rows)

    if not active_manual_groups and not batch_rows:
        st.error("请至少填写一组手动地址，或上传一份批量地址文件。")
    elif uploaded_batch_file is not None and batch_file_error:
        st.error(batch_file_error)
    elif manual_validation_errors or batch_validation_errors or cross_source_duplicate_errors:
        st.error("检测到数据问题，已阻止生成，请先修正以下内容：")
        for message in manual_validation_errors + batch_validation_errors + cross_source_duplicate_errors:
            st.write(f"- {message}")
    else:
        try:
            parsed_address_groups = [
                {
                    "address_info": split_address(str(address_group["address_text"])),
                    "order_groups": address_group["order_groups"],
                }
                for address_group in active_manual_groups
            ]
            if batch_rows:
                parsed_address_groups.extend(
                    [
                        {
                            "address_info": split_address(row["address_text"]),
                            "order_groups": [
                                {
                                    "external_order_no": row["external_order_no"],
                                    "line_items": batch_shared_line_items,
                                }
                            ],
                        }
                        for row in batch_rows
                    ]
                )

            output_file = build_output_file(
                template_path=TEMPLATE_PATH,
                address_groups=parsed_address_groups,
                customer_code=selected_customer["code"],
                customer_name=selected_customer["name"],
                payment_date=now_shanghai().date(),
            )
            file_name = f"样品寄送_{now_shanghai().strftime('%Y%m%d_%H%M%S')}.xlsx"
            saved_path = save_output_file(BASE_DIR, output_file, file_name)
            total_orders = sum(len(address_group["order_groups"]) for address_group in parsed_address_groups)
            total_items = sum(
                len(order_group["line_items"])
                for address_group in parsed_address_groups
                for order_group in address_group["order_groups"]
            )
            st.success(
                f"Excel 已生成，共写入 {len(parsed_address_groups)} 个地址、{total_orders} 个外部单号、{total_items} 行商品，并已保存到 output 文件夹：{saved_path.name}"
            )
            st.download_button(
                label="下载生成的 Excel",
                data=output_file.getvalue(),
                file_name=file_name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        except ValueError as exc:
            st.error(str(exc))
